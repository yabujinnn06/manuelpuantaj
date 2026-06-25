# -*- coding: utf-8 -*-
"""WhatsApp dokumunden cikarilan Mayis 2026 puantaj kontrol tablosu."""
from datetime import date, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

OUT = r"C:\Users\canor\OneDrive\Masaüstü\puantaj\puantaj_mayis_2026_kontrol.xlsx"

# Sadece cikis bildiren (giris hic yazilmamis) teknik personel
tech_cikis = {
    "Hasan Tontur": {1:"18:00",2:"21:00",4:"19:00",5:"20:00",6:"19:30",7:"19:30",
                     8:"19:00",9:"17:00",11:"18:00",12:"17:30",13:"19:00",14:"20:30",
                     15:"19:30",16:"18:30",18:"19:00"},
    "Ata Türkbey": {1:"18:30",2:"19:00",4:"18:30",5:"20:00",6:"20:10",7:"18:30",
                    8:"19:30",9:"17:00",11:"18:30",12:"19:00",13:"19:30",14:"19:30",
                    15:"20:00",16:"18:00",18:"18:40"},
}

# Giris + cikis bildirenler. "IZIN" = izinli. 3'lu tuple -> 3. eleman not.
gc = {
    "Uğur Ertürk": {1:"IZIN",2:("09:30","14:00"),4:("09:50","21:50"),5:("09:50","20:10"),
        6:("09:50","18:30"),7:("09:30","18:30"),8:("09:50","18:30"),9:("09:50","15:30"),
        11:("09:10","18:30"),12:("09:40","20:50"),13:("09:55","18:30"),14:("09:20","21:30"),
        15:("09:25","18:30"),16:("09:30","21:30"),18:("09:30","19:00")},
    "Ercüment Çalışkan": {1:("09:30","19:30"),2:("09:30","14:00"),4:("09:30","21:30"),
        5:("09:30","20:50"),6:("09:30","19:15"),7:("09:30","18:30"),8:("09:30","20:30"),
        9:("09:30","14:30"),11:("09:30","18:30"),12:("09:30","18:30"),13:("09:30","18:30"),
        14:("09:30","20:30"),15:("09:30","20:50"),16:("09:30","20:30"),
        18:("09:30","20:10","Demo nedeniyle mesajda 12:00 giris notu var")},
    "Eda Nur Yılmaz": {1:("10:00","22:00","Mesajda tarih 01.03.2026 yazilmis, 01.05 kabul edildi"),
        2:("10:00","18:00","Mesajda tarih 02.05.206 yazilmis"),3:("14:00","22:00"),
        4:("14:00","22:00"),5:("14:00","22:00"),6:("10:00","18:00","Mesajda tarih 06,05.2026"),
        7:"IZIN",8:("14:00","22:00"),9:("10:00","18:00"),10:("14:00","22:00"),
        11:("10:00","18:00"),12:("14:00","22:00"),13:("10:00","22:00"),14:("10:00","18:00"),
        15:"IZIN",16:("14:00","22:00"),17:("10:00","18:00"),18:("10:00","22:00")},
    "Kübracan Gündoğdu": {1:("10:00","22:00"),2:("10:00","18:00"),3:("14:00","22:00"),
        4:"IZIN",5:("14:00","22:00"),6:("10:00","18:00"),7:("14:00","22:00"),
        8:("10:00","22:00"),9:("10:00","18:00"),10:("14:00","22:00"),11:("10:00","18:00"),
        12:("10:00","22:00"),13:("10:00","18:00"),14:"IZIN",15:("14:00","22:00"),
        16:("14:00","22:00"),17:("10:00","18:00"),18:"IZIN"},
    "Başak Çelik": {1:"IZIN",2:("14:00","22:00"),3:("10:00","18:00"),4:("10:00","22:00"),
        5:("10:00","18:00"),6:("14:00","22:00"),7:("10:00","18:00"),8:"IZIN",
        9:("14:00","22:00"),10:("10:00","18:00"),11:("14:00","22:00"),12:"IZIN",
        13:("14:00","22:00"),14:("10:00","22:00"),15:("10:00","18:00"),16:("10:00","18:00"),
        17:("14:00","22:00"),18:("10:00","22:00")},
    "Leyla Karayağız": {1:"IZIN",2:("14:00","22:00"),3:("10:00","18:00"),4:("10:00","18:00"),
        5:"IZIN",6:("14:00","22:00"),7:("10:00","22:00"),8:("10:00","18:00"),
        9:("14:00","22:00"),10:("10:00","18:00"),11:("14:00","22:00"),12:("10:00","18:00"),
        13:"IZIN",14:("14:00","22:00"),15:("10:00","22:00"),16:("10:00","18:00"),
        17:("14:00","22:00"),18:"IZIN"},
}

def hours(g, c):
    fmt = "%H:%M"
    t1 = datetime.strptime(g, fmt); t2 = datetime.strptime(c, fmt)
    d = (t2 - t1).total_seconds() / 3600.0
    if d < 0:
        d += 24
    return round(d, 2)

rows = []  # (date, person, giris, cikis, sure, durum, not)
for day in range(1, 19):
    d = date(2026, 5, day)
    ds = d.strftime("%Y-%m-%d")
    for person, dd in tech_cikis.items():
        if day in dd:
            c = dd[day]
            rows.append((ds, person, "09:30", c, hours("09:30", c), "NORMAL",
                         "Giris bildirilmemis, 09:30 varsayildi"))
    for person, dd in gc.items():
        if day not in dd:
            continue
        v = dd[day]
        if v == "IZIN":
            rows.append((ds, person, "", "", "", "IZIN", ""))
        else:
            g, c = v[0], v[1]
            note = v[2] if len(v) > 2 else ""
            rows.append((ds, person, g, c, hours(g, c), "NORMAL", note))

rows.sort(key=lambda r: (r[0], r[1]))

wb = Workbook()
ws = wb.active
ws.title = "Mayis 2026"
hdr = ["Tarih", "Personel", "Giris", "Cikis", "Sure (s)", "Durum", "Not"]
ws.append(hdr)
hf = Font(bold=True, color="FFFFFF")
hfill = PatternFill("solid", fgColor="305496")
for i, _ in enumerate(hdr, 1):
    cl = ws.cell(row=1, column=i)
    cl.font = hf; cl.fill = hfill; cl.alignment = Alignment(horizontal="center")

red = PatternFill("solid", fgColor="F8CBAD")
yellow = PatternFill("solid", fgColor="FFF2CC")
for r in rows:
    ws.append(r)
    rr = ws.max_row
    if r[5] == "GİRİŞ YOK":
        for c in range(1, 8):
            ws.cell(row=rr, column=c).fill = red
    elif r[5] == "IZIN":
        for c in range(1, 8):
            ws.cell(row=rr, column=c).fill = yellow
    elif r[6]:
        ws.cell(row=rr, column=7).fill = yellow

widths = [12, 20, 8, 8, 9, 12, 48]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
ws.freeze_panes = "A2"

# Sorunlar sayfasi
ws2 = wb.create_sheet("Sorunlar ve Eksikler")
issues = [
    ["Konu", "Aciklama"],
    ["Kapsam", "WhatsApp dokumu 18.05.2026'da bitiyor. 19-31 Mayis verisi YOK."],
    ["Burhan Findikzade", "Mayis 2026'da HIC kayit yok (gruptaki son kaydi Nisan). Tel: +90 505 074 24 61"],
    ["Hasan Tontur", "Giris saati hic bildirilmemis. Talimat geregi tum gunler icin giris 09:30 varsayildi, sure ona gore hesaplandi."],
    ["Ata Türkbey", "Giris saati hic bildirilmemis. Talimat geregi tum gunler icin giris 09:30 varsayildi, sure ona gore hesaplandi."],
    ["Pazar gunleri", "03, 10, 17 Mayis Pazar. Teknik/montaj ekibi (Hasan, Ata, Ugur, Ercument) bu gunlerde kayit girmemis (muhtemelen tatil). Stant ekibi calismis."],
    ["Ugur Ertürk", "01.05 izinli. 03/10/17 Mayis (Pazar) kayit yok."],
    ["Ercüment Çalışkan", "18.05 mesajinda 'demo nedeniyle 12:00 giris verildi' notu var; bildirilen giris 09:30."],
    ["Eda Nur Yilmaz", "Tarih yazim hatalari: 01.03.2026 (->01.05), 02.05.206 (->02.05), 06,05.2026 (->06.05). Icerikten duzeltildi."],
    ["Dilara Akbaba", "01.03.2026'da gruptan cikarilmis; Mayis'ta zaten kayit yok. Kapsam disi."],
    ["Genel", "Stant ekibi giris+cikis duzenli veriyor. Teknik ekip cogunlukla sadece cikis veriyor; giris saatleri sistemde varsayilan/elle tamamlanmali."],
]
for row in issues:
    ws2.append(row)
for c in range(1, 3):
    cl = ws2.cell(row=1, column=c); cl.font = hf; cl.fill = hfill
ws2.column_dimensions["A"].width = 22
ws2.column_dimensions["B"].width = 95
for rr in range(1, ws2.max_row + 1):
    ws2.cell(row=rr, column=2).alignment = Alignment(wrap_text=True, vertical="top")

wb.save(OUT)
print("Yazildi:", OUT)
print("Toplam kayit:", len(rows))
