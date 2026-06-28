"""WhatsApp parser, preview Excel ve bulk apply icin testler.

SAMPLE gercek "Rainwater puantaj" grubu formatini taklit eder: Turkce saat
damgali baslik, gonderen != calisan, cok-satirli mesaj, tek cikis saati,
Full/izin/pazar mesaisi, cok-gunlu mesaj.
"""

import os
import sys

import pytest

HERE = os.path.dirname(__file__)
APP_DIR = os.path.abspath(os.path.join(HERE, "..", "..", "..", ".."))
HARNESS_DIR = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
for p in (APP_DIR, HARNESS_DIR):
    if p not in sys.path:
        sys.path.insert(0, p)


from cli_anything.puantaj import whatsapp  # noqa: E402


SAMPLE = """\
6.01.2026 öğleden sonra 4:50 - Hasan Teknik: 05.01.2026 çıkış saati 23:00
Hasan TONTUR
6.01.2026 öğleden sonra 4:14 - +90 537 732 05 42: Ata Türkbey
05.01.2026 çıkış saati 19:00
28.03.2026 akşam 7:10 - Ercüment Abi Rainwater: 28.03.2026
Ercüment ÇALIŞKAN
GİRİŞ : 09:30
ÇIKIŞ : 19:10
28.03.2026 öğleden sonra 5:46 - +90 532 621 64 06: 28.03.2026 Uğur Ertürk 09:28 giriş 17:30 çıkış
29.03.2026 gece 10:05 - +90 534 771 93 17: 29.03.2026
Eda Nur Yılmaz
Giriş 10:00
Çıkış 18:00
17.01.2026 öğleden sonra 4:19 - Başak Çelik Rain: Başak Çelik
17.01.2026
Full
18.02.2026 akşam 7:35 - +90 534 771 93 17: 18.02.2026
Eda Nur
Haftalık izin
15.02.2026 öğleden sonra 3:45 - Ercüment Abi Rainwater: Ercument çalışkan  15.02.2026 ÇIKIŞ: PAZAR MESAİSİ
"""

# Cok-gunlu tek mesaj (Hasan)
MULTI_DAY = """\
7.01.2026 öğleden önce 11:34 - Hasan Teknik: Hasan TONTUR
01.01.26
YILBAŞI

05.01.26
08:00 GİRİŞ
23:00 çıkış

06.01.26
08:00. GİRİŞ
5.30 çıkış
"""


def test_turkce_header_and_basic_parse():
    entries = whatsapp.parse_text(SAMPLE, region="Ankara")
    # En az 8 kayit cikmali
    assert len(entries) >= 8
    names = {e.employee_name_raw for e in entries}
    assert any("Hasan" in n for n in names)
    assert any("Ercüment" in n or "Ercument" in n for n in names)


def test_sender_phone_uses_body_name():
    """Gonderen telefon numarasi oldugunda isim mesaj govdesinden alinmali."""
    entries = whatsapp.parse_text(SAMPLE, region="Ankara")
    ata = next(e for e in entries if "Ata" in e.employee_name_raw)
    assert "Türkbey" in ata.employee_name_raw or "Turkbey" in ata.employee_name_raw
    assert ata.work_date == "2026-01-05"
    assert ata.end_time == "19:00"


def test_single_exit_time():
    """Sadece cikis saati yazilan kayitlarda end_time dolu, status Calisti."""
    entries = whatsapp.parse_text(SAMPLE, region="Ankara")
    hasan = next(e for e in entries if "Hasan" in e.employee_name_raw)
    assert hasan.status == "Calisti"
    assert hasan.end_time == "23:00"
    assert hasan.work_date == "2026-01-05"


def test_multiline_giris_cikis():
    """Cok satirli 'GIRIS .. / CIKIS ..' dogru ayrismali (karismamali)."""
    entries = whatsapp.parse_text(SAMPLE, region="Ankara")
    erc = next(e for e in entries
               if e.work_date == "2026-03-28" and "Ercüment" in e.employee_name_raw)
    assert erc.start_time == "09:30"
    assert erc.end_time == "19:10"


def test_inline_giris_cikis_same_line():
    """'09:28 giris 17:30 cikis' tek satirda dogru ayrismali."""
    entries = whatsapp.parse_text(SAMPLE, region="Ankara")
    ugur = next(e for e in entries if "Uğur" in e.employee_name_raw)
    assert ugur.start_time == "09:28"
    assert ugur.end_time == "17:30"


def test_full_expands_to_shift():
    entries = whatsapp.parse_text(SAMPLE, region="Ankara")
    basak = next(e for e in entries if "Başak" in e.employee_name_raw)
    assert basak.start_time == "10:00"
    assert basak.end_time == "22:00"
    assert basak.status == "Calisti"


def test_leave_detected():
    entries = whatsapp.parse_text(SAMPLE, region="Ankara")
    izin = [e for e in entries if e.status == "Izinli"]
    assert izin and any("Eda" in e.employee_name_raw for e in izin)


def test_sunday_work_flag():
    entries = whatsapp.parse_text(SAMPLE, region="Ankara")
    pazar = [e for e in entries if e.is_sunday_work]
    assert pazar and pazar[0].work_date == "2026-02-15"


def test_multi_day_message_keeps_name():
    """Cok-gunlu mesajda isim her gune tasinmali; not satiri isim sayilmamali."""
    entries = whatsapp.parse_text(MULTI_DAY, region="Ankara")
    assert all("Hasan" in e.employee_name_raw for e in entries), \
        [e.employee_name_raw for e in entries]
    dates = {e.work_date for e in entries}
    assert "2026-01-01" in dates  # YILBASI
    assert "2026-01-05" in dates
    j5 = next(e for e in entries if e.work_date == "2026-01-05")
    assert j5.start_time == "08:00"
    assert j5.end_time == "23:00"


def test_sender_based_identity_no_phantom_names():
    """Ayni gonderenin tum mesajlari tek kisidir; sohbet/not satirlari
    uydurma calisan uretmez."""
    text = (
        "7.01.2026 öğleden önce 11:34 - Hasan Teknik: Hasan TONTUR\n"
        "01.01.26\n"
        "YILBAŞI\n"
        "02.01.26\n"
        "Giriş 08:30\n"
        "Stan kuruldu ertesi sabaha kadar\n"
        "12.01.2026 akşam 4:43 - Hasan Teknik: Teşekkürler. Sabah 06:00 /22:00 arası. Pazar mesai si\n"
        "13.01.2026 akşam 6:46 - Hasan Teknik: 13.01.26\n"
        "Hasan tontur\n"
        "Giriş 7.30\n"
        "Çıkış 18:45\n"
    )
    entries = whatsapp.parse_text(text, region="Ankara")
    names = {e.employee_name_raw for e in entries}
    # Tek kisi olmali; "Stan kuruldu", "Tesekkurler" gibi isimler OLMAMALI
    assert len(names) == 1, names
    only = next(iter(names))
    assert "Hasan" in only
    assert "Stan" not in only and "esekkur" not in only.lower()


def test_phone_sender_grouped_as_one_person():
    text = (
        "13.01.2026 akşamüstü 6:40 - +90 537 732 05 42: 13.01.26 Ata Türkbey\n"
        "Çıkış saati 17:30\n"
        "14.01.2026 akşam 7:05 - +90 537 732 05 42: 14.01.2026\n"
        "Çıkış 19:04\n"
        "Ata Türkbey\n"
    )
    entries = whatsapp.parse_text(text, region="Ankara")
    names = {e.employee_name_raw for e in entries}
    assert len(names) == 1
    assert "Ata" in next(iter(names))


def test_year_typo_corrected():
    text = ("6.01.2026 öğleden sonra 5:31 - Hasan Teknik: 06.01.2016 çıkış 17.30\n"
            "Hasan TONTUR\n")
    entries = whatsapp.parse_text(text, region="Ankara")
    assert entries
    assert entries[0].work_date == "2026-01-06"
    assert any("duzeltildi" in w.lower() or "düzeltildi" in w.lower()
               for w in entries[0].warnings)


def test_space_separated_year_not_time():
    # "07.01 2026 Cikis: 02:00" -> yil bosluklu yazilmis; 2026 saat (20:26) sanilmamali
    text = ("7.01.2026 sabah 10:51 - Ercüment Abi Rainwater: 07.01 2026 Çıkış: 02:00\n")
    entries = whatsapp.parse_text(text, region="Ankara")
    assert entries
    assert entries[0].end_time == "02:00"


def test_example_template_message_skipped():
    # "*Ornek: ... Cikis Saati: 17:30*" sablon mesaji kayit uretmemeli
    text = ("6.01.2026 öğleden sonra 12:54 - Hüseyincan Orman: "
            "*Örnek: Hüseyincan Orman 6.01.2026 Çıkış Saati: 17:30*\n")
    entries = whatsapp.parse_text(text, region="Ankara")
    workers = [e for e in entries if not e.is_non_worker]
    assert workers == []


def test_admin_group_creator_flagged_non_worker():
    # Grubu kuran kisi (admin) puantaj yazsa bile calisan sayilmaz.
    text = (
        "6.01.2026 öğleden önce 11:13 - ‎Altan Akbaş Abi \"Ankara rainwater puantaj\" grubunu oluşturdu\n"
        "19.01.2026 öğleden sonra 4:20 - Altan Akbaş Abi: Stanttaki arkadaşlar full yazdıkları zaman 10:00-22:00\n"
    )
    entries = whatsapp.parse_text(text, region="Ankara")
    # Altan'in uretebildigi tum kayitlar non_worker isaretli olmali
    assert all(e.is_non_worker for e in entries)


def test_ik_sender_flagged_non_worker():
    text = ("20.05.2026 öğleden sonra 5:41 - Begüm Hanim Rainwater İk: "
            "Merhabalar, rapor durumu hakkında bilgi.\n")
    entries = whatsapp.parse_text(text, region="Ankara")
    assert all(e.is_non_worker for e in entries)


def test_real_worker_not_flagged():
    # Govdede kendi adiyla puantaj yazan gercek calisan non_worker OLMAMALI.
    text = ("6.01.2026 öğleden sonra 4:14 - +90 537 732 05 42: Ata Türkbey\n"
            "05.01.2026 çıkış saati 19:00\n")
    entries = whatsapp.parse_text(text, region="Ankara")
    assert entries
    assert not any(e.is_non_worker for e in entries)


def test_shift_code_resolved():
    # "14-10" vardiya kodu -> 14:00-22:00 (saat 14:00 giris, 22:00 cikis)
    assert whatsapp._resolve_shift_code(14, 10) == ("14:00", "22:00")
    assert whatsapp._resolve_shift_code(2, 10) == ("14:00", "22:00")
    assert whatsapp._resolve_shift_code(10, 10) == ("10:00", "22:00")
    assert whatsapp._resolve_shift_code(10, 18) == ("10:00", "18:00")
    assert whatsapp._resolve_shift_code(7, 30) is None  # 30 dk -> vardiya degil


def test_exit_before_entry_treated_as_pm():
    # "Giris 14:00 / Cikis 10:00" -> cikis aksam 10 = 22:00 (gunduz vardiyasi)
    text = ("31.01.2026 gece 10:00 - +90 534 771 93 17: 31.01.2026\n"
            "Eda Nur\n"
            "Giriş: 14:00\n"
            "Çıkış 10:00\n")
    entries = whatsapp.parse_text(text, region="Ankara")
    rec = next(e for e in entries if e.status == "Calisti")
    assert rec.start_time == "14:00"
    assert rec.end_time == "22:00"


def test_equal_entry_exit_clears_start():
    text = ("30.03.2026 akşam 7:00 - +90 532 621 64 06: 30.03.2026 Uğur 09:50 giriş 09:50 çıkış\n")
    entries = whatsapp.parse_text(text, region="Ankara")
    rec = next(e for e in entries if e.status == "Calisti")
    # Esit saatte giris bosaltilir, cikis korunur
    assert rec.end_time == "09:50"
    assert rec.start_time is None


def test_early_morning_exit_uses_afternoon_default():
    # Cikis 02:00 (gece yarisi sonrasi) -> giris varsayilani ogleden sonra
    start, assumed = whatsapp.department_default_start("Stant", "02:00")
    assert start == "16:00" and assumed is True


def test_clock_normalization():
    assert whatsapp.normalize_clock("19.30") == "19:30"
    assert whatsapp.normalize_clock("0800") == "08:00"
    assert whatsapp.normalize_clock("5.30") == "05:30"
    assert whatsapp.normalize_clock("08;30") == "08:30"
    assert whatsapp.normalize_clock("07/30") == "07:30"
    assert whatsapp.normalize_clock("9:5") is None


def test_system_messages_skipped():
    text = ("6.01.2026 öğleden önce 11:13 - Altan Akbaş Abi sizi ekledi\n"
            "6.01.2026 öğleden sonra 4:14 - +90 537 732 05 42: Ata Türkbey\n"
            "05.01.2026 çıkış saati 19:00\n")
    entries = whatsapp.parse_text(text, region="Ankara")
    assert len(entries) == 1
    assert "Ata" in entries[0].employee_name_raw


def test_match_employees_by_full_name():
    employees = [
        (1, "Hasan Tontur", "111", "TEKNIK", "Tek", "Ankara"),
        (2, "Ercument Caliskan", "222", "OFIS", "Uzman", "Ankara"),
    ]
    entries = whatsapp.parse_text(SAMPLE, region="Ankara")
    whatsapp.match_employees(entries, employees, region="Ankara")
    hasan = next(e for e in entries if "Hasan" in e.employee_name_raw)
    assert hasan.employee_id == 1
    assert hasan.confidence >= 0.6


def test_no_match_warns():
    employees = [(1, "Hasan Tontur", "", "", "", "Ankara")]
    entries = whatsapp.parse_text(SAMPLE, region="Ankara")
    whatsapp.match_employees(entries, employees, region="Ankara")
    ata = next(e for e in entries if "Ata" in e.employee_name_raw)
    assert ata.employee_id is None
    assert any("eslesmedi" in w.lower() for w in ata.warnings)


def test_apply_shift_defaults_fills_start():
    """Giris yoksa departman varsayilani uygulanir, isaretlenir."""
    employees = [(1, "Hasan Tontur", "", "TEKNIK", "", "Ankara")]
    entries = whatsapp.parse_text(SAMPLE, region="Ankara")
    whatsapp.match_employees(entries, employees, region="Ankara")
    whatsapp.apply_shift_defaults(entries)
    hasan = next(e for e in entries
                 if e.employee_id == 1 and e.status == "Calisti" and e.end_time)
    assert hasan.start_time is not None
    assert hasan.start_assumed is True


# ---------------------------------------------------------------- preview + apply
@pytest.fixture
def temp_db(monkeypatch, tmp_path):
    monkeypatch.setenv("APPDATA", str(tmp_path))
    for mod in ("puantaj_db",):
        if mod in sys.modules:
            del sys.modules[mod]
    import puantaj_db as db
    db.init_db()
    db.add_employee("Hasan Tontur", "111", "TEKNIK", "Tek", "Ankara")
    db.add_employee("Ata Turkbey", "222", "TEKNIK", "Tek", "Ankara")
    db.add_employee("Ercument Caliskan", "333", "OFIS", "Uzman", "Ankara")
    db.add_employee("Ugur Erturk", "444", "LOJISTIK", "Sofor", "Ankara")
    db.add_employee("Eda Nur Yilmaz", "555", "STANT", "Sat", "Ankara")
    db.add_employee("Basak Celik", "666", "STANT", "Sat", "Ankara")
    return db


def _payload(db):
    entries = whatsapp.parse_text(SAMPLE, region="Ankara")
    whatsapp.match_employees(entries, db.list_employees(), region="Ankara")
    whatsapp.apply_shift_defaults(entries)
    return whatsapp.entries_to_dicts(entries)


def _empmap(db):
    return {int(r[0]): {"full_name": r[1], "department": r[3], "region": r[5]}
            for r in db.list_employees()}


def test_preview_excel_sheets(temp_db, tmp_path):
    from cli_anything.puantaj import preview_xlsx
    out = tmp_path / "preview.xlsx"
    path = preview_xlsx.build_preview(
        str(out), _payload(temp_db), employees_by_id=_empmap(temp_db),
        settings=temp_db.get_all_settings(),
        shift_templates=temp_db.list_shift_templates(),
    )
    assert os.path.isfile(path) and os.path.getsize(path) > 1000
    from openpyxl import load_workbook
    wb = load_workbook(path)
    for sheet in ("Ozet", "Calisan Matrisi", "Calisan Analizi", "Detay",
                  "Gunluk Ozet", "Uyarilar"):
        assert sheet in wb.sheetnames


def test_preview_per_employee_sheets(temp_db, tmp_path):
    from cli_anything.puantaj import preview_xlsx
    out = tmp_path / "preview.xlsx"
    preview_xlsx.build_preview(
        str(out), _payload(temp_db), employees_by_id=_empmap(temp_db),
        settings=temp_db.get_all_settings(),
        shift_templates=temp_db.list_shift_templates(),
    )
    from openpyxl import load_workbook
    wb = load_workbook(str(out))
    # Hasan icin ayri sekme olmali
    assert any("Hasan" in s for s in wb.sheetnames)


def test_apply_writes_records(temp_db):
    from cli_anything.puantaj import bulk
    result = bulk.apply_entries(temp_db, _payload(temp_db), default_region="Ankara")
    assert result.timesheets_added >= 4
    ts = temp_db.list_timesheets()
    att = temp_db.list_attendance_records()
    assert ts and att
