import os
import logging
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas as pdf_canvas
    from reportlab.lib.utils import ImageReader
except Exception:
    pdf_canvas = None

from calc import calc_day_hours, is_sunday_non_stand

logger = logging.getLogger("rainstaff")


HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
TITLE_FILL = PatternFill("solid", fgColor="EAF2FB")
THIN = Side(border_style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def calc_sunday_separate_hours(work_date, department, is_special, worked_hours):
    if is_special:
        return 0.0
    try:
        if is_sunday_non_stand(work_date, department):
            return max(0.0, float(worked_hours))
    except Exception:
        return 0.0
    return 0.0


def unpack_timesheet_record(record):
    """
    Desteklenen formatlar:
    - Yeni: (id, emp_id, name, department, work_date, start, end, break, is_special, notes, region)
    - Eski: (id, emp_id, name, work_date, start, end, break, is_special, notes, region)
    """
    if len(record) >= 11:
        (
            _,
            emp_id,
            name,
            department,
            work_date,
            start_time,
            end_time,
            break_minutes,
            is_special,
            notes,
            region,
            *_rest,
        ) = record
    elif len(record) == 10:
        (
            _,
            emp_id,
            name,
            work_date,
            start_time,
            end_time,
            break_minutes,
            is_special,
            notes,
            region,
        ) = record
        department = ""
    else:
        raise ValueError(f"Desteklenmeyen puantaj kayit formati: {record}")
    return (
        emp_id,
        name,
        department,
        work_date,
        start_time,
        end_time,
        break_minutes,
        is_special,
        notes,
        region,
    )


def _new_total_row(name):
    return {
        "name": name,
        "worked": 0.0,
        "scheduled": 0.0,
        "overtime": 0.0,
        "night": 0.0,
        "overnight": 0.0,
        "special_normal": 0.0,
        "special_overtime": 0.0,
        "special_night": 0.0,
        "sunday_hours": 0.0,
        "sunday_days": set(),
        "worked_days": set(),
        "leave_days": set(),
        "attendance_days": set(),
    }


def _parse_iso_date(value):
    try:
        return datetime.strptime(str(value), "%Y-%m-%d").date()
    except Exception:
        return None


def _iter_date_keys(start_date, end_date, clamp_start=None, clamp_end=None):
    start_dt = _parse_iso_date(start_date)
    end_dt = _parse_iso_date(end_date)
    if not start_dt or not end_dt:
        return
    if end_dt < start_dt:
        return
    if clamp_start and start_dt < clamp_start:
        start_dt = clamp_start
    if clamp_end and end_dt > clamp_end:
        end_dt = clamp_end
    if end_dt < start_dt:
        return
    day = start_dt
    while day <= end_dt:
        yield day.strftime("%Y-%m-%d")
        day += timedelta(days=1)


def export_report(
    output_path,
    records,
    settings,
    date_range_text,
    attendance_records=None,
    leave_records=None,
    start_date=None,
    end_date=None,
):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rainstaff"

    company_name = settings.get("company_name", "")
    report_title = settings.get("report_title", "Rainstaff Puantaj ve Mesai Raporu")
    logo_path = settings.get("logo_path", "")

    row = 1
    header_cols = 17
    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 18

    for r in range(1, 4):
        for c in range(1, header_cols + 1):
            ws.cell(row=r, column=c).fill = TITLE_FILL

    if logo_path and os.path.isfile(logo_path):
        try:
            img = Image(logo_path)
            img.width = 110
            img.height = 60
            ws.add_image(img, "A1")
        except Exception as e:
            logger.warning("Logo yukleme basarısız (%s): %s", logo_path, str(e))

    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=header_cols)
    title_cell = ws.cell(row=row, column=2, value=company_name)
    title_cell.font = Font(size=14, bold=True)
    title_cell.alignment = Alignment(vertical="center", horizontal="left")
    row += 1
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=header_cols)
    subtitle_cell = ws.cell(row=row, column=2, value=report_title)
    subtitle_cell.font = Font(size=11, bold=True)
    subtitle_cell.alignment = Alignment(vertical="center", horizontal="left")
    row += 1
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=header_cols)
    date_cell = ws.cell(row=row, column=2, value=date_range_text or "Tarih Araligi: -")
    date_cell.font = Font(size=10)
    date_cell.alignment = Alignment(vertical="center", horizontal="left")
    row += 2

    headers = [
        "Calisan",
        "Bolge",
        "Tarih",
        "Giris",
        "Cikis",
        "Mola (dk)",
        "Calisilan (s)",
        "Plan (s)",
        "Fazla Mesai (s)",
        "Gece (s)",
        "Geceye Tasan (s)",
        "Ozel Gun",
        "Ozel Gun Normal (s)",
        "Ozel Gun Fazla (s)",
        "Ozel Gun Gece (s)",
        "Pazar Mesaisi (s)",
        "Not",
    ]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER
    row += 1

    totals = {}
    special_records = []
    overnight_records = []
    leave_section_records = []
    attendance_records = attendance_records or []
    leave_records = leave_records or []
    clamp_start = _parse_iso_date(start_date)
    clamp_end = _parse_iso_date(end_date)
    for record in records:
        (
            emp_id,
            name,
            department,
            work_date,
            start_time,
            end_time,
            break_minutes,
            is_special,
            notes,
            region,
        ) = unpack_timesheet_record(record)
        (
            worked,
            scheduled,
            overtime,
            night_hours,
            overnight_hours,
            special_normal,
            special_overtime,
            special_night,
        ) = calc_day_hours(
            work_date,
            start_time,
            end_time,
            break_minutes,
            settings,
            is_special,
            department,
        )
        ws.cell(row=row, column=1, value=name).border = BORDER
        ws.cell(row=row, column=2, value=region or "").border = BORDER
        ws.cell(row=row, column=3, value=work_date).border = BORDER
        ws.cell(row=row, column=4, value=start_time).border = BORDER
        ws.cell(row=row, column=5, value=end_time).border = BORDER
        ws.cell(row=row, column=6, value=break_minutes).border = BORDER
        ws.cell(row=row, column=7, value=worked).border = BORDER
        ws.cell(row=row, column=8, value=scheduled).border = BORDER
        ws.cell(row=row, column=9, value=overtime).border = BORDER
        ws.cell(row=row, column=10, value=night_hours).border = BORDER
        ws.cell(row=row, column=11, value=overnight_hours).border = BORDER
        ws.cell(row=row, column=12, value="Evet" if is_special else "Hayir").border = BORDER
        ws.cell(row=row, column=13, value=special_normal).border = BORDER
        ws.cell(row=row, column=14, value=special_overtime).border = BORDER
        ws.cell(row=row, column=15, value=special_night).border = BORDER
        sunday_hours = calc_sunday_separate_hours(work_date, department, is_special, worked)
        ws.cell(row=row, column=16, value=sunday_hours).border = BORDER
        ws.cell(row=row, column=17, value=notes or "").border = BORDER

        if emp_id not in totals:
            totals[emp_id] = _new_total_row(name)
        elif not totals[emp_id].get("name") and name:
            totals[emp_id]["name"] = name
        work_date_key = str(work_date or "").strip()
        totals[emp_id]["worked"] += worked
        totals[emp_id]["scheduled"] += scheduled
        totals[emp_id]["overtime"] += overtime
        totals[emp_id]["night"] += night_hours
        totals[emp_id]["overnight"] += overnight_hours
        totals[emp_id]["special_normal"] += special_normal
        totals[emp_id]["special_overtime"] += special_overtime
        totals[emp_id]["special_night"] += special_night
        totals[emp_id]["sunday_hours"] += sunday_hours
        if work_date_key:
            totals[emp_id]["worked_days"].add(work_date_key)
            if sunday_hours > 0:
                totals[emp_id]["sunday_days"].add(work_date_key)
        if is_special:
            special_records.append((name, work_date, special_normal, special_overtime, special_night))
        if overnight_hours > 0:
            overnight_records.append((name, work_date, overnight_hours))
        row += 1

    for attendance in attendance_records:
        if len(attendance) < 5:
            continue
        try:
            _att_id, emp_id, name, work_date, status, *_rest = attendance
        except Exception:
            continue
        if emp_id not in totals:
            totals[emp_id] = _new_total_row(name)
        elif not totals[emp_id].get("name") and name:
            totals[emp_id]["name"] = name
        work_date_key = str(work_date or "").strip()
        if not work_date_key:
            continue
        totals[emp_id]["attendance_days"].add(work_date_key)
        status_text = str(status or "").strip().lower()
        if status_text == "calisti":
            totals[emp_id]["worked_days"].add(work_date_key)
        elif status_text == "izinli":
            totals[emp_id]["leave_days"].add(work_date_key)
            totals[emp_id]["worked_days"].discard(work_date_key)

    for leave in leave_records:
        if len(leave) < 10:
            continue
        try:
            (
                _leave_id,
                emp_id,
                name,
                leave_start,
                leave_end,
                leave_type,
                reason,
                _doc_no,
                _doc_path,
                leave_status,
                *_rest,
            ) = leave
        except Exception:
            continue

        status_text = str(leave_status or "").strip()
        in_range_days = list(_iter_date_keys(leave_start, leave_end, clamp_start, clamp_end))
        leave_section_records.append(
            (
                name,
                leave_start,
                leave_end,
                len(in_range_days),
                leave_type or "",
                status_text or "",
                reason or "",
            )
        )
        if status_text.lower() not in {"onayli", "approved", ""}:
            continue
        if emp_id not in totals:
            totals[emp_id] = _new_total_row(name)
        elif not totals[emp_id].get("name") and name:
            totals[emp_id]["name"] = name
        for day_key in in_range_days:
            totals[emp_id]["leave_days"].add(day_key)
            totals[emp_id]["worked_days"].discard(day_key)

    row += 1
    ws.cell(row=row, column=1, value="Ozet").font = Font(bold=True)
    row += 1
    summary_headers = [
        "Calisan",
        "Toplam Calisilan (s)",
        "Toplam Plan (s)",
        "Toplam Fazla Mesai (s)",
        "Toplam Gece (s)",
        "Toplam Geceye Tasan (s)",
        "Pazar Gun",
        "Pazar Mesaisi (s)",
        "Calisilan Gun",
        "Izinli Gun",
    ]
    for col, header in enumerate(summary_headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
    row += 1

    for _, data in sorted(totals.items(), key=lambda x: x[1]["name"]):
        ws.cell(row=row, column=1, value=data["name"]).border = BORDER
        ws.cell(row=row, column=2, value=round(data["worked"], 2)).border = BORDER
        ws.cell(row=row, column=3, value=round(data["scheduled"], 2)).border = BORDER
        ws.cell(row=row, column=4, value=round(data["overtime"], 2)).border = BORDER
        ws.cell(row=row, column=5, value=round(data["night"], 2)).border = BORDER
        ws.cell(row=row, column=6, value=round(data["overnight"], 2)).border = BORDER
        ws.cell(row=row, column=7, value=len(data["sunday_days"])).border = BORDER
        ws.cell(row=row, column=8, value=round(data["sunday_hours"], 2)).border = BORDER
        ws.cell(row=row, column=9, value=len(data["worked_days"])).border = BORDER
        ws.cell(row=row, column=10, value=len(data["leave_days"])).border = BORDER
        row += 1

    row += 2
    ws.cell(row=row, column=1, value="Ozel Gun Calismalari").font = Font(bold=True)
    row += 1
    special_headers = [
        "Calisan",
        "Tarih",
        "Ozel Gun Normal (s)",
        "Ozel Gun Fazla (s)",
        "Ozel Gun Gece (s)",
    ]
    for col, header in enumerate(special_headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
    row += 1
    if special_records:
        for name, work_date, spec_norm, spec_ot, spec_night in special_records:
            ws.cell(row=row, column=1, value=name).border = BORDER
            ws.cell(row=row, column=2, value=work_date).border = BORDER
            ws.cell(row=row, column=3, value=spec_norm).border = BORDER
            ws.cell(row=row, column=4, value=spec_ot).border = BORDER
            ws.cell(row=row, column=5, value=spec_night).border = BORDER
            row += 1
    else:
        ws.cell(row=row, column=1, value="Kayit yok").border = BORDER
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Geceye Tasan Mesailer").font = Font(bold=True)
    row += 1
    overnight_headers = ["Calisan", "Tarih", "Geceye Tasan (s)"]
    for col, header in enumerate(overnight_headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
    row += 1
    if overnight_records:
        for name, work_date, overnight in overnight_records:
            ws.cell(row=row, column=1, value=name).border = BORDER
            ws.cell(row=row, column=2, value=work_date).border = BORDER
            ws.cell(row=row, column=3, value=overnight).border = BORDER
            row += 1
    else:
        ws.cell(row=row, column=1, value="Kayit yok").border = BORDER
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Izin Kayitlari").font = Font(bold=True)
    row += 1
    leave_headers = ["Calisan", "Baslangic", "Bitis", "Gun", "Izin Tipi", "Durum", "Neden"]
    for col, header in enumerate(leave_headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
    row += 1
    if leave_section_records:
        for name, leave_start, leave_end, leave_days, leave_type, leave_status, leave_reason in leave_section_records:
            ws.cell(row=row, column=1, value=name).border = BORDER
            ws.cell(row=row, column=2, value=leave_start).border = BORDER
            ws.cell(row=row, column=3, value=leave_end).border = BORDER
            ws.cell(row=row, column=4, value=leave_days).border = BORDER
            ws.cell(row=row, column=5, value=leave_type).border = BORDER
            ws.cell(row=row, column=6, value=leave_status).border = BORDER
            ws.cell(row=row, column=7, value=leave_reason).border = BORDER
            row += 1
    else:
        ws.cell(row=row, column=1, value="Kayit yok").border = BORDER
        row += 1

    for col in range(1, 18):
        ws.column_dimensions[chr(64 + col)].width = 16

    ws.freeze_panes = "A6"
    wb.save(output_path)


def export_leave_form(
    output_path,
    employee_name,
    identity_no,
    department,
    title,
    start_date,
    end_date,
    leave_type,
    reason,
    company_name,
    document_no,
    logo_path="",
):
    wb = Workbook()
    ws = wb.active
    ws.title = "Izin Formu"

    header_font = Font(size=14, bold=True)
    sub_font = Font(size=11, bold=True)
    label_font = Font(size=10, bold=True, color="2F2F2F")
    value_font = Font(size=10)
    section_fill = PatternFill("solid", fgColor="EAF2FB")
    light_fill = PatternFill("solid", fgColor="F7F9FC")

    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 18

    for r in range(1, 4):
        for c in range(1, 7):
            ws.cell(row=r, column=c).fill = section_fill

    if logo_path and os.path.isfile(logo_path):
        try:
            img = Image(logo_path)
            img.width = 110
            img.height = 60
            ws.add_image(img, "A1")
        except Exception as e:
            logger.warning("Logo yukleme basarısız (%s): %s", logo_path, str(e))

    ws.merge_cells("B1:F1")
    ws["B1"] = company_name or " "
    ws["B1"].font = header_font
    ws["B1"].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("B2:F2")
    ws["B2"] = "IZIN FORMU"
    ws["B2"].font = sub_font
    ws["B2"].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("B3:F3")
    ws["B3"] = f"Belge No: {document_no or '-'}"
    ws["B3"].font = Font(size=10)
    ws["B3"].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("A5:F5")
    ws["A5"] = "Personel Bilgisi"
    ws["A5"].font = sub_font
    ws["A5"].fill = section_fill
    ws["A5"].alignment = Alignment(horizontal="left")

    ws["A6"] = "Calisan Adi"
    ws["A6"].font = label_font
    ws["B6"] = employee_name
    ws["B6"].font = value_font
    ws["D6"] = "Unvan"
    ws["D6"].font = label_font
    ws["E6"] = title or "-"
    ws["E6"].font = value_font

    ws["A7"] = "TC Kimlik"
    ws["A7"].font = label_font
    ws["B7"] = identity_no or "-"
    ws["B7"].font = value_font
    ws["D7"] = "Departman"
    ws["D7"].font = label_font
    ws["E7"] = department or "-"
    ws["E7"].font = value_font

    ws.merge_cells("A9:F9")
    ws["A9"] = "Izin Bilgisi"
    ws["A9"].font = sub_font
    ws["A9"].fill = section_fill
    ws["A9"].alignment = Alignment(horizontal="left")

    ws["A10"] = "Izin Tipi"
    ws["A10"].font = label_font
    ws["B10"] = leave_type or "-"
    ws["B10"].font = value_font
    ws["D10"] = "Toplam Gun"
    ws["D10"].font = label_font

    try:
        start_dt = datetime.strptime(start_date, "%Y-%m-%d")
        end_dt = datetime.strptime(end_date, "%Y-%m-%d")
        days = (end_dt.date() - start_dt.date()).days + 1
    except Exception:
        days = ""
    ws["E10"] = days
    ws["E10"].font = value_font

    ws["A11"] = "Baslangic"
    ws["A11"].font = label_font
    ws["B11"] = start_date
    ws["B11"].font = value_font
    ws["D11"] = "Bitis"
    ws["D11"].font = label_font
    ws["E11"] = end_date
    ws["E11"].font = value_font

    ws.merge_cells("A13:F13")
    ws["A13"] = "Aciklama"
    ws["A13"].font = sub_font
    ws["A13"].fill = section_fill
    ws["A13"].alignment = Alignment(horizontal="left")

    ws.merge_cells("A14:F16")
    ws["A14"] = reason or "-"
    ws["A14"].font = value_font
    ws["A14"].fill = light_fill
    ws["A14"].alignment = Alignment(wrap_text=True, vertical="top")

    ws["A18"] = "Calisan Imza"
    ws["A18"].font = label_font
    ws["D18"] = "Yonetici Imza"
    ws["D18"].font = label_font

    for col in range(1, 7):
        ws.column_dimensions[chr(64 + col)].width = 18

    for row in range(5, 19):
        for col in range(1, 7):
            ws.cell(row=row, column=col).border = BORDER

    wb.save(output_path)


def export_report_pdf(output_path, records, settings, date_range_text, max_rows=200):
    if pdf_canvas is None:
        raise RuntimeError("PDF icin reportlab kurulu degil.")

    c = pdf_canvas.Canvas(output_path, pagesize=A4)
    width, height = A4
    y = height - 40
    company_name = settings.get("company_name", "Rainstaff")
    report_title = settings.get("report_title", "Puantaj ve Mesai Raporu")
    logo_path = settings.get("logo_path", "")

    if logo_path and os.path.isfile(logo_path):
        try:
            img = ImageReader(logo_path)
            c.drawImage(img, 40, height - 80, width=80, height=40, preserveAspectRatio=True, mask="auto")
        except Exception:
            pass

    c.setFont("Helvetica-Bold", 13)
    c.drawString(130, y, company_name)
    y -= 18
    c.setFont("Helvetica-Bold", 10)
    c.drawString(130, y, report_title)
    y -= 16
    c.setFont("Helvetica", 9)
    c.drawString(130, y, date_range_text or "Tarih Araligi: -")
    y -= 22

    headers = ["Calisan", "Tarih", "Giris", "Cikis", "Calisilan", "Fazla"]
    col_x = [40, 220, 300, 350, 410, 470]
    c.setFont("Helvetica-Bold", 8)
    for h, x in zip(headers, col_x):
        c.drawString(x, y, h)
    y -= 10
    c.line(40, y, width - 40, y)
    y -= 12

    c.setFont("Helvetica", 8)
    count = 0
    for record in records:
        (
            _emp_id,
            name,
            department,
            work_date,
            start_time,
            end_time,
            break_minutes,
            is_special,
            notes,
            region,
        ) = unpack_timesheet_record(record)
        if count >= max_rows:
            break
        try:
            worked, _scheduled, overtime, _night, _overnight, _s1, _s2, _s3 = calc_day_hours(
                work_date,
                start_time,
                end_time,
                break_minutes,
                settings,
                is_special,
                department,
            )
        except Exception:
            worked = overtime = 0
        c.drawString(col_x[0], y, str(name)[:24])
        c.drawString(col_x[1], y, str(work_date))
        c.drawString(col_x[2], y, str(start_time))
        c.drawString(col_x[3], y, str(end_time))
        c.drawString(col_x[4], y, f"{worked:.2f}")
        c.drawString(col_x[5], y, f"{overtime:.2f}")
        y -= 12
        count += 1
        if y < 60:
            c.showPage()
            y = height - 40
            c.setFont("Helvetica", 8)

    c.save()


def export_vehicle_weekly_report(
    output_path,
    plate,
    week_start,
    prev_week,
    checklist,
    prev_results,
    current_results,
    current_km,
    prev_km,
    vehicle_row,
    current_fault,
    prev_fault,
    service_visits,
):
    wb = Workbook()
    ws = wb.active
    ws.title = "Arac Kontrol"

    ws.cell(row=1, column=1, value="Arac Haftalik Kontrol Raporu").font = Font(size=12, bold=True)
    ws.cell(row=2, column=1, value=f"Plaka: {plate}")
    ws.cell(row=3, column=1, value=f"Hafta: {week_start}")
    ws.cell(row=4, column=1, value=f"Onceki Hafta: {prev_week or '-'}")
    ws.cell(row=5, column=1, value=f"KM (Bu Hafta): {current_km or '-'}")
    ws.cell(row=6, column=1, value=f"KM (Onceki): {prev_km or '-'}")

    if vehicle_row:
        (
            _vid,
            _plate,
            brand,
            model,
            year,
            km,
            inspection_date,
            insurance_date,
            maintenance_date,
            oil_change_date,
            oil_change_km,
            oil_interval_km,
            _notes,
            _region,
        ) = vehicle_row
        today = datetime.now().date()
        if inspection_date:
            insp_dt = datetime.strptime(inspection_date, "%Y-%m-%d").date()
            diff = (insp_dt - today).days
            insp_text = f"{inspection_date} ({diff} gun)" if diff >= 0 else f"{inspection_date} ({abs(diff)} gun gecikme)"
        else:
            insp_text = "-"
        ws.cell(row=7, column=1, value=f"Muayene: {insp_text}")
        ws.cell(row=8, column=1, value=f"Son Bakim: {maintenance_date or '-'}")
        ws.cell(row=9, column=1, value=f"Son Yag Degisimi: {oil_change_date or '-'}")
        interval_km = oil_interval_km or 14000
        if interval_km and oil_change_km is not None and current_km:
            remaining = interval_km - (current_km - oil_change_km)
            oil_text = "Geldi" if remaining <= 0 else f"{remaining} km"
        else:
            oil_text = "-"
        ws.cell(row=10, column=1, value=f"Yag Periyodu: {interval_km or '-'} km, Kalan: {oil_text}")

    row = 12
    ws.cell(row=row, column=1, value="Ariza Bilgisi").font = Font(bold=True)
    row += 1
    ws.cell(row=row, column=1, value=f"Bu Hafta Ariza: {current_fault.get('title') or '-'}")
    row += 1
    ws.cell(row=row, column=1, value=f"Bu Hafta Durum: {current_fault.get('status') or '-'}")
    row += 1
    ws.cell(
        row=row,
        column=1,
        value=f"Sanayiye Gitti: {'Evet' if current_fault.get('service') else 'Hayir'}",
    )
    row += 1
    ws.cell(row=row, column=1, value=f"Onceki Hafta Ariza: {prev_fault.get('title') or '-'}")
    row += 1
    ws.cell(row=row, column=1, value=f"Onceki Hafta Durum: {prev_fault.get('status') or '-'}")
    row += 1
    ws.cell(
        row=row,
        column=1,
        value=f"Onceki Hafta Sanayi: {'Evet' if prev_fault.get('service') else 'Hayir'}",
    )
    row += 2

    ws.cell(row=row, column=1, value="Sanayi Kayitlari (Hafta)").font = Font(bold=True)
    row += 1
    svc_headers = ["Gidis", "Donus", "Neden", "Masraf", "Not"]
    for col, header in enumerate(svc_headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
    row += 1
    if service_visits:
        for visit in service_visits:
            _sid, _vid, _plate, _fid, _title, start_date, end_date, reason, cost, notes, _region = visit
            ws.cell(row=row, column=1, value=start_date or "-").border = BORDER
            ws.cell(row=row, column=2, value=end_date or "Sanayide").border = BORDER
            ws.cell(row=row, column=3, value=reason or "-").border = BORDER
            ws.cell(row=row, column=4, value=cost if cost is not None else "-").border = BORDER
            ws.cell(row=row, column=5, value=notes or "-").border = BORDER
            row += 1
    else:
        ws.cell(row=row, column=1, value="Kayit yok").border = BORDER
        row += 1
    row += 1

    headers = ["Kontrol", "Onceki Hafta", "Bu Hafta", "Durum"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
    row += 1

    def normalize_status(value):
        text = str(value or "").strip()
        mapping = {
            "OK": "Olumlu",
            "Issue": "Olumsuz",
            "NA": "Bilinmiyor",
            "Olumlu": "Olumlu",
            "Olumsuz": "Olumsuz",
            "Bilinmiyor": "Bilinmiyor",
        }
        return mapping.get(text, text or "-")

    for item_key, label in checklist:
        prev_status = normalize_status(prev_results.get(item_key, "-"))
        curr_status = normalize_status(current_results.get(item_key, "-"))
        if prev_status == curr_status:
            change = "Ayni"
        elif prev_status == "Olumsuz" and curr_status == "Olumsuz":
            change = "Tekrar eden sorun"
        elif curr_status == "Olumsuz" and prev_status != "Olumsuz":
            change = "Kotulesti"
        elif prev_status == "Olumsuz" and curr_status != "Olumsuz":
            change = "Iyilesti"
        else:
            change = "Degisti"

        ws.cell(row=row, column=1, value=label).border = BORDER
        ws.cell(row=row, column=2, value=prev_status).border = BORDER
        ws.cell(row=row, column=3, value=curr_status).border = BORDER
        ws.cell(row=row, column=4, value=change).border = BORDER
        row += 1

    for col in range(1, 6):
        ws.column_dimensions[chr(64 + col)].width = 22

    wb.save(output_path)


def export_vehicle_card_report(output_path, plate, vehicle_row, inspections, faults, services):
    wb = Workbook()
    ws = wb.active
    ws.title = "Arac Karti"

    ws.cell(row=1, column=1, value="Arac Karti").font = Font(size=12, bold=True)
    ws.cell(row=2, column=1, value=f"Plaka: {plate}")

    if vehicle_row:
        (
            _vid,
            _plate,
            brand,
            model,
            year,
            km,
            inspection_date,
            insurance_date,
            maintenance_date,
            oil_change_date,
            oil_change_km,
            oil_interval_km,
            notes,
            _region,
        ) = vehicle_row
        ws.cell(row=3, column=1, value=f"Marka/Model: {brand} {model}")
        ws.cell(row=4, column=1, value=f"Yil: {year}")
        ws.cell(row=5, column=1, value=f"KM: {km or '-'}")
        ws.cell(row=6, column=1, value=f"Muayene: {inspection_date or '-'}")
        ws.cell(row=7, column=1, value=f"Sigorta: {insurance_date or '-'}")
        ws.cell(row=8, column=1, value=f"Bakim: {maintenance_date or '-'}")
        ws.cell(row=9, column=1, value=f"Yag Degisim: {oil_change_date or '-'}")
        ws.cell(row=10, column=1, value=f"Yag KM: {oil_change_km or '-'}")
        ws.cell(row=11, column=1, value=f"Yag Periyot: {(oil_interval_km or 14000)} km")
        if notes:
            ws.cell(row=12, column=1, value=f"Not: {notes}")

    fault_title_map = {row[0]: row[3] for row in faults} if faults else {}

    row = 14
    ws.cell(row=row, column=1, value="Ariza Kayitlari").font = Font(bold=True)
    row += 1
    headers = ["Baslik", "Durum", "Acilis", "Kapanis", "Aciklama"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
    row += 1
    if faults:
        for fault in faults:
            _fid, _vid, _plate, title, desc, opened_date, closed_date, status, _region = fault
            ws.cell(row=row, column=1, value=title or "-").border = BORDER
            ws.cell(row=row, column=2, value=status or "-").border = BORDER
            ws.cell(row=row, column=3, value=opened_date or "-").border = BORDER
            ws.cell(row=row, column=4, value=closed_date or "-").border = BORDER
            ws.cell(row=row, column=5, value=desc or "-").border = BORDER
            row += 1
    else:
        ws.cell(row=row, column=1, value="Kayit yok").border = BORDER
        row += 1

    row += 2
    ws.cell(row=row, column=1, value="Sanayi Kayitlari").font = Font(bold=True)
    row += 1
    headers = ["Ariza", "Gidis", "Donus", "Masraf", "Neden", "Not"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
    row += 1
    if services:
        for visit in services:
            _sid, _vid, _plate, _fid, title, start_date, end_date, reason, cost, notes, _region = visit
            ws.cell(row=row, column=1, value=title or "-").border = BORDER
            ws.cell(row=row, column=2, value=start_date or "-").border = BORDER
            ws.cell(row=row, column=3, value=end_date or "Sanayide").border = BORDER
            ws.cell(row=row, column=4, value=cost if cost is not None else "-").border = BORDER
            ws.cell(row=row, column=5, value=reason or "-").border = BORDER
            ws.cell(row=row, column=6, value=notes or "-").border = BORDER
            row += 1
    else:
        ws.cell(row=row, column=1, value="Kayit yok").border = BORDER
        row += 1

    row += 2
    ws.cell(row=row, column=1, value="Kontroller").font = Font(bold=True)
    row += 1
    headers = ["Tarih", "Hafta", "Surucu", "KM", "Ariza", "Durum", "Sanayi", "Not"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
    row += 1
    if inspections:
        for row_data in inspections:
            (
                _iid,
                _veh_id,
                _plate,
                _driver_id,
                driver_name,
                inspect_date,
                week_start,
                km_val,
                note_val,
                fault_id,
                fault_status,
                service_visit,
            ) = row_data
            ws.cell(row=row, column=1, value=inspect_date or "-").border = BORDER
            ws.cell(row=row, column=2, value=week_start or "-").border = BORDER
            ws.cell(row=row, column=3, value=driver_name or "-").border = BORDER
            ws.cell(row=row, column=4, value=km_val or "-").border = BORDER
            ws.cell(row=row, column=5, value=fault_title_map.get(fault_id, "-")).border = BORDER
            ws.cell(row=row, column=6, value=fault_status or "-").border = BORDER
            ws.cell(row=row, column=7, value="Evet" if service_visit else "Hayir").border = BORDER
            ws.cell(row=row, column=8, value=note_val or "-").border = BORDER
            row += 1
    else:
        ws.cell(row=row, column=1, value="Kayit yok").border = BORDER
        row += 1

    for col in range(1, 9):
        ws.column_dimensions[chr(64 + col)].width = 20

    wb.save(output_path)
